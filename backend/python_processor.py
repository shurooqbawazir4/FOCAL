#!/usr/bin/env python3
"""
FlashCard Generator - Python File Processor

This module handles advanced file processing and AI-powered flashcard generation.

Features:
- Multi-format file support (PDF, PowerPoint, Images)
- Intelligent text extraction and cleaning
- OCR processing for images
- OpenAI integration for context-aware flashcard generation
- Robust error handling and logging

Author: FlashCard Generator Team
Version: 1.0.0
"""

import sys
import json
import os
from pathlib import Path
import tempfile
import logging

# Import libraries for different file types
try:
    import PyPDF2
    import pdfplumber
except ImportError:
    PyPDF2 = None
    pdfplumber = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    from PIL import Image
    import pytesseract
except ImportError:
    Image = None
    pytesseract = None

try:
    import openai
    # Try both old and new OpenAI API formats
    try:
        from openai import OpenAI
        NEW_OPENAI = True
    except ImportError:
        NEW_OPENAI = False
except ImportError:
    openai = None
    OpenAI = None
    NEW_OPENAI = False

# Set up logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class FileProcessor:
    def __init__(self, openai_api_key=None):
        """Initialize the file processor with OpenAI API key."""
        self.openai_client = None
        self.openai_api_key = openai_api_key
        
        if openai_api_key and openai:
            try:
                if NEW_OPENAI:
                    # Use new OpenAI client
                    self.openai_client = OpenAI(api_key=openai_api_key)
                else:
                    # Use old OpenAI API
                    openai.api_key = openai_api_key
                    self.openai_client = "legacy"
                logger.info("OpenAI client initialized successfully")
            except Exception as e:
                logger.error(f"Failed to initialize OpenAI client: {e}")
                # Try legacy approach
                try:
                    openai.api_key = openai_api_key
                    self.openai_client = "legacy"
                    logger.info("Using legacy OpenAI API")
                except Exception as e2:
                    logger.error(f"Legacy OpenAI initialization also failed: {e2}")

    def extract_text_from_pdf(self, file_path):
        """Extract text from PDF using multiple methods."""
        text_content = ""
        
        # Try pdfplumber first (better for complex layouts)
        if pdfplumber:
            try:
                with pdfplumber.open(file_path) as pdf:
                    for page in pdf.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text_content += page_text + "\n"
                logger.info(f"pdfplumber extracted {len(text_content)} characters")
            except Exception as e:
                logger.warning(f"pdfplumber failed: {e}")
        
        # Fallback to PyPDF2 if pdfplumber didn't work well
        if len(text_content.strip()) < 50 and PyPDF2:
            try:
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    for page in pdf_reader.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text_content += page_text + "\n"
                logger.info(f"PyPDF2 extracted {len(text_content)} characters")
            except Exception as e:
                logger.warning(f"PyPDF2 failed: {e}")
        
        return text_content.strip()

    def extract_text_from_powerpoint(self, file_path):
        """Extract text from PowerPoint presentations."""
        text_content = ""
        
        # Check if it's an old .ppt file
        if file_path.lower().endswith('.ppt'):
            logger.warning("Old .ppt format detected. Converting to text using alternative method...")
            
            # For old .ppt files, try to read as binary and extract readable text
            try:
                with open(file_path, 'rb') as file:
                    content = file.read()
                    # Convert to string and extract readable text
                    text = content.decode('latin-1', errors='ignore')
                    
                    # Extract readable ASCII text chunks
                    import re
                    text_chunks = re.findall(r'[A-Za-z0-9\s\.,!?;:\-()]{10,}', text)
                    
                    # Filter and clean text chunks
                    meaningful_chunks = []
                    for chunk in text_chunks:
                        chunk = chunk.strip()
                        if (len(chunk) > 15 and 
                            chunk.count(' ') > 2 and  # Has multiple words
                            not chunk.startswith('Microsoft') and  # Skip metadata
                            'Times New Roman' not in chunk and
                            'Arial' not in chunk):
                            meaningful_chunks.append(chunk)
                    
                    if meaningful_chunks:
                        text_content = '\n'.join(meaningful_chunks[:20])  # Limit to first 20 chunks
                        logger.info(f"Extracted {len(meaningful_chunks)} text chunks from .ppt file")
                    
                    if len(text_content.strip()) < 50:
                        raise ValueError("Could not extract meaningful text from .ppt file")
                        
            except Exception as e:
                logger.error(f"Failed to process .ppt file: {e}")
                raise ValueError("Unable to process old .ppt format. Please convert to .pptx or PDF format for better results.")
        
        else:
            # Handle .pptx files with python-pptx
            if not Presentation:
                raise ImportError("python-pptx not installed. Install with: pip install python-pptx")
            
            try:
                prs = Presentation(file_path)
                
                for slide_num, slide in enumerate(prs.slides, 1):
                    slide_text = f"\n--- Slide {slide_num} ---\n"
                    
                    # Extract text from all shapes in the slide
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            slide_text += shape.text.strip() + "\n"
                        
                        # Handle tables
                        if shape.has_table:
                            table = shape.table
                            for row in table.rows:
                                row_text = []
                                for cell in row.cells:
                                    if cell.text.strip():
                                        row_text.append(cell.text.strip())
                                if row_text:
                                    slide_text += " | ".join(row_text) + "\n"
                    
                    # Only add slide if it has meaningful content
                    if len(slide_text.strip()) > 20:
                        text_content += slide_text
                
                logger.info(f"PowerPoint extracted {len(text_content)} characters from {len(prs.slides)} slides")
                
            except Exception as e:
                logger.error(f"Failed to extract text from PowerPoint: {e}")
                raise
        
        if not text_content.strip():
            raise ValueError("No text content could be extracted from the PowerPoint file")
            
        return text_content.strip()

    def extract_text_from_image(self, file_path):
        """Extract text from images using OCR."""
        if not (Image and pytesseract):
            raise ImportError("PIL and pytesseract not installed. Install with: pip install Pillow pytesseract")
        
        try:
            # Open and preprocess the image
            image = Image.open(file_path)
            
            # Convert to grayscale for better OCR
            if image.mode != 'L':
                image = image.convert('L')
            
            # Extract text using OCR
            text = pytesseract.image_to_string(image, lang='eng')
            
            logger.info(f"OCR extracted {len(text)} characters from image")
            return text.strip()
            
        except Exception as e:
            logger.error(f"Failed to extract text from image: {e}")
            raise

    def process_file(self, file_path, file_type):
        """Process a file and extract its text content."""
        logger.info(f"Processing file: {file_path} (type: {file_type})")
        
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
        
        text_content = ""
        
        # Determine file type and extract accordingly
        if file_type.lower() in ['pdf', 'application/pdf']:
            text_content = self.extract_text_from_pdf(file_path)
        elif file_type.lower() in ['ppt', 'pptx', 'application/vnd.ms-powerpoint', 
                                  'application/vnd.openxmlformats-officedocument.presentationml.presentation']:
            text_content = self.extract_text_from_powerpoint(file_path)
        elif file_type.lower() in ['jpg', 'jpeg', 'png', 'gif', 'bmp', 'tiff'] or 'image' in file_type.lower():
            text_content = self.extract_text_from_image(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")
        
        if not text_content or len(text_content.strip()) < 10:
            raise ValueError("No meaningful text content could be extracted from the file")
        
        return text_content

    def generate_flashcards(self, text_content, num_cards=None):
        """Generate flashcards from text content using OpenAI."""
        if not self.openai_client:
            raise ValueError("OpenAI client not initialized. Please provide a valid API key.")
        
        # Determine number of cards based on content length
        if num_cards is None:
            content_length = len(text_content)
            if content_length < 500:
                num_cards = "3-5"
            elif content_length < 2000:
                num_cards = "5-8"
            elif content_length < 5000:
                num_cards = "8-12"
            else:
                num_cards = "10-15"
        
        prompt = f"""
You are an expert educational content analyzer. Create high-quality flashcards from the following content.

CONTENT TO ANALYZE:
{text_content[:6000]}  # Limit to avoid token limits

INSTRUCTIONS:
1. Create {num_cards} flashcards that capture the most important concepts
2. Focus on key facts, definitions, processes, and important details
3. Make questions specific and clear
4. Provide comprehensive but concise answers
5. Ensure questions test understanding, not just memorization
6. Use varied question types (what, why, how, when, etc.)

OUTPUT FORMAT:
Return ONLY a valid JSON array with objects containing "question" and "answer" fields.
Do not include any other text or explanations.

Example format:
[
  {{"question": "What is the main concept discussed?", "answer": "The main concept is..."}},
  {{"question": "Why is this important?", "answer": "This is important because..."}}
]
"""

        try:
            logger.info("Generating flashcards with OpenAI...")
            
            if self.openai_client == "legacy":
                # Use legacy OpenAI API
                response = openai.ChatCompletion.create(
                    model="gpt-3.5-turbo",
                    messages=[
                        {
                            "role": "system",
                            "content": "You are an expert educator who creates high-quality study flashcards. Always respond with valid JSON only."
                        },
                        {
                            "role": "user",
                            "content": prompt
                        }
                    ],
                    max_tokens=2000,
                    temperature=0.7
                )
                response_text = response.choices[0].message.content.strip()
            else:
                # Use new OpenAI client
                response = self.openai_client.chat.completions.create(
                    model="gpt-3.5-turbo",
                    messages=[
                        {
                            "role": "system",
                            "content": "You are an expert educator who creates high-quality study flashcards. Always respond with valid JSON only."
                        },
                        {
                            "role": "user",
                            "content": prompt
                        }
                    ],
                    max_tokens=2000,
                    temperature=0.7
                )
                response_text = response.choices[0].message.content.strip()
            
            logger.info(f"OpenAI response length: {len(response_text)}")
            
            # Parse JSON response
            try:
                flashcards = json.loads(response_text)
            except json.JSONDecodeError:
                # Try to extract JSON from response if it's wrapped in text
                import re
                json_match = re.search(r'\[.*\]', response_text, re.DOTALL)
                if json_match:
                    flashcards = json.loads(json_match.group())
                else:
                    raise ValueError("Could not parse JSON from OpenAI response")
            
            # Validate flashcards
            if not isinstance(flashcards, list):
                raise ValueError("Response is not a list of flashcards")
            
            valid_flashcards = []
            for card in flashcards:
                if (isinstance(card, dict) and 
                    'question' in card and 'answer' in card and
                    card['question'].strip() and card['answer'].strip()):
                    valid_flashcards.append(card)
            
            if not valid_flashcards:
                raise ValueError("No valid flashcards generated")
            
            logger.info(f"Generated {len(valid_flashcards)} valid flashcards")
            return valid_flashcards
            
        except Exception as e:
            logger.error(f"Failed to generate flashcards: {e}")
            raise

def main():
    """Main function for command-line usage."""
    if len(sys.argv) < 4:
        print("Usage: python3 python_processor.py <file_path> <file_type> <openai_api_key>")
        sys.exit(1)
    
    file_path = sys.argv[1]
    file_type = sys.argv[2]
    openai_api_key = sys.argv[3]
    
    try:
        processor = FileProcessor(openai_api_key)
        
        # Extract text from file
        text_content = processor.process_file(file_path, file_type)
        logger.info(f"Extracted {len(text_content)} characters of text")
        
        # Generate flashcards
        flashcards = processor.generate_flashcards(text_content)
        
        # Output result as JSON
        result = {
            "success": True,
            "text_length": len(text_content),
            "flashcards": flashcards,
            "text_preview": text_content[:300] + "..." if len(text_content) > 300 else text_content
        }
        
        print(json.dumps(result, indent=2))
        
    except Exception as e:
        error_result = {
            "success": False,
            "error": str(e)
        }
        print(json.dumps(error_result, indent=2))
        sys.exit(1)

if __name__ == "__main__":
    main()
